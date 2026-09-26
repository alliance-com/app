from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
from datetime import datetime
import os

DOSSIER_SORTIE = "bons_paiement"
os.makedirs(DOSSIER_SORTIE, exist_ok=True)

# Couleurs modernes
BLEU_FONCE = RGBColor(15, 40, 80)
BLEU_MOYEN = RGBColor(30, 90, 160)
GRIS_TEXTE = RGBColor(50, 50, 50)
GRIS_CLAIR = RGBColor(120, 120, 120)
VERT_OK = RGBColor(20, 120, 70)

def set_run_font(run, size=10, bold=False, color=None, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color

def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)

def set_cell_margins(cell, top=40, bottom=40, left=60, right=60):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

def add_bottom_border(paragraph, color="1E5AA0", size="12"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def ajouter_filigrane_centre(doc, reference):
    """
    Vrai filigrane au centre de la page (derrière le texte).
    STE NAP SARL + N° référence + date du jour.
    N'ajoute aucune page.
    """
    date_str = datetime.now().strftime("%d/%m/%Y")
    texte = f"STE NAP SARL  •  {reference}  •  {date_str}"

    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    for p in header.paragraphs:
        p.clear()
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()

    r = OxmlElement("w:r")
    paragraph._p.append(r)

    rPr = OxmlElement("w:rPr")
    noProof = OxmlElement("w:noProof")
    rPr.append(noProof)
    r.append(rPr)

    pict = OxmlElement("w:pict")
    r.append(pict)

    shape_xml = f'''
    <v:shape xmlns:v="urn:schemas-microsoft-com:vml"
             xmlns:o="urn:schemas-microsoft-com:office:office"
             xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
             id="FiligraneNAP"
             type="#_x0000_t136"
             style="position:absolute;margin-left:0;margin-top:0;width:500pt;height:60pt;z-index:-251658240;mso-position-horizontal:center;mso-position-horizontal-relative:page;mso-position-vertical:center;mso-position-vertical-relative:page;"
             fillcolor="#C8C8C8"
             stroked="f">
      <v:fill opacity=".40"/>
      <v:textbox inset="0,0,0,0">
        <w:txbxContent>
          <w:p>
            <w:pPr>
              <w:jc w:val="center"/>
            </w:pPr>
            <w:r>
              <w:rPr>
                <w:b/>
                <w:sz w:val="28"/>
                <w:szCs w:val="28"/>
                <w:color w:val="A0A0A0"/>
              </w:rPr>
              <w:t xml:space="preserve">{texte}</w:t>
            </w:r>
          </w:p>
        </w:txbxContent>
      </v:textbox>
    </v:shape>
    '''
    shape = etree.fromstring(shape_xml)
    pict.append(shape)

def generer_bon_paiement(data):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(1.0)
        section.bottom_margin = Cm(1.0)
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # ----- FILIGRANE au centre de la page -----
    reference = str(data.get("No_transaction", ""))
    ajouter_filigrane_centre(doc, reference)

    # ========== BANDEAU ENTÊTE ==========
    table_header = doc.add_table(rows=1, cols=1)
    table_header.autofit = True
    cell_h = table_header.rows[0].cells[0]
    shade_cell(cell_h, "0F2850")
    set_cell_margins(cell_h, top=80, bottom=80, left=100, right=100)

    p = cell_h.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("STE NAP SARL")
    set_run_font(run, size=18, bold=True, color=RGBColor(255, 255, 255))

    p2 = cell_h.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(
        "PLOT NO 14 GLODJIGBE ECONOMIC ZONE • REPUBLIC OF BENIN\n"
        "RCCM RB/COT/21 B 30710  •  IFU 3202113334014  •  TEL +229 01 94 94 04 14 / 01 97 97 45 06"
    )
    set_run_font(run2, size=8, color=RGBColor(200, 220, 255))

    doc.add_paragraph("")

    # ========== TITRE ==========
    titre = doc.add_paragraph()
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = titre.add_run("BON DE PAIEMENT")
    set_run_font(rt, size=16, bold=True, color=BLEU_MOYEN)
    add_bottom_border(titre, color="1E5AA0", size="18")

    doc.add_paragraph("")

    # ========== BLOC INFOS (2 colonnes) ==========
    table_info = doc.add_table(rows=5, cols=4)
    table_info.autofit = True

    gauche = [
        ("Date", str(data.get("Date", "")).replace(" 00:00:00", "")),
        ("N° de transaction", str(data.get("No_transaction", ""))),
        ("Type d'opération", str(data.get("Type_operation", ""))),
        ("Produit", str(data.get("Produit", ""))),
    ]
    droite = [
        ("Prestataire / Fournisseur", str(data.get("Prestataire", ""))),
        ("Nom du percepteur", str(data.get("Beneficiaire", ""))),
        ("N° Pièce d'identité", str(data.get("Piece_identite", ""))),
        ("Date d'expiration", str(data.get("Date_expiration", "")).replace(" 00:00:00", "")),
        ("Adresse / Téléphone", str(data.get("Telephone", "") or "—")),
    ]

    for i in range(5):
        if i < len(gauche):
            label, val = gauche[i]
            c0 = table_info.rows[i].cells[0]
            c1 = table_info.rows[i].cells[1]
            c0.text = ""
            c1.text = ""
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(label)
            set_run_font(r0, size=8, bold=True, color=GRIS_CLAIR)
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(val)
            set_run_font(r1, size=10, bold=True, color=BLEU_FONCE)
            set_cell_margins(c0, top=30, bottom=30)
            set_cell_margins(c1, top=30, bottom=30)

        label, val = droite[i]
        c2 = table_info.rows[i].cells[2]
        c3 = table_info.rows[i].cells[3]
        c2.text = ""
        c3.text = ""
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(label)
        set_run_font(r2, size=8, bold=True, color=GRIS_CLAIR)
        p3 = c3.paragraphs[0]
        r3 = p3.add_run(val)
        set_run_font(r3, size=10, bold=True, color=BLEU_FONCE)
        set_cell_margins(c2, top=30, bottom=30)
        set_cell_margins(c3, top=30, bottom=30)

    doc.add_paragraph("")

    # ========== DÉTAILS DES OPÉRATIONS ==========
    p_sec = doc.add_paragraph()
    rs = p_sec.add_run("DÉTAILS DES OPÉRATIONS")
    set_run_font(rs, size=10, bold=True, color=BLEU_MOYEN)
    add_bottom_border(p_sec, color="1E5AA0", size="8")

    table_det = doc.add_table(rows=3, cols=4)
    table_det.style = "Table Grid"
    table_det.autofit = True

    headers = ["Réf. Doc", "Description / Objet du paiement", "Montant", "Observation"]
    for i, h in enumerate(headers):
        cell = table_det.rows[0].cells[i]
        cell.text = ""
        shade_cell(cell, "1E5AA0")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, size=9, bold=True, color=RGBColor(255, 255, 255))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell, top=50, bottom=50)

    vals = [
        str(data.get("Ref_doc", "—") or "—"),
        str(data.get("Description", "")),
        f"XOF {float(data.get('Montant', 0)):,.0f}".replace(",", " "),
        str(data.get("Observation", "")),
    ]
    for i, v in enumerate(vals):
        cell = table_det.rows[1].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(v)
        set_run_font(r, size=9, color=GRIS_TEXTE)
        if i in (2, 3):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell, top=60, bottom=60)

    table_det.rows[2].cells[0].merge(table_det.rows[2].cells[1])
    cell_total_label = table_det.rows[2].cells[0]
    cell_total_val = table_det.rows[2].cells[2]
    cell_obs = table_det.rows[2].cells[3]

    shade_cell(cell_total_label, "F0F5FA")
    shade_cell(cell_total_val, "F0F5FA")
    shade_cell(cell_obs, "F0F5FA")

    cell_total_label.text = ""
    p = cell_total_label.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("MONTANT TOTAL  ")
    set_run_font(r, size=10, bold=True, color=BLEU_FONCE)

    cell_total_val.text = ""
    p = cell_total_val.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"XOF {float(data.get('Montant', 0)):,.0f}".replace(",", " "))
    set_run_font(r, size=11, bold=True, color=VERT_OK)

    cell_obs.text = ""

    doc.add_paragraph("")

    # ========== MONTANT EN LETTRES ==========
    table_lettres = doc.add_table(rows=1, cols=1)
    cell_l = table_lettres.rows[0].cells[0]
    shade_cell(cell_l, "F7F9FC")
    set_cell_margins(cell_l, top=60, bottom=60, left=80, right=80)
    p = cell_l.paragraphs[0]
    r1 = p.add_run("Montant en lettres :  ")
    set_run_font(r1, size=9, bold=True, color=GRIS_CLAIR)
    r2 = p.add_run(str(data.get("Montant_lettres", "")))
    set_run_font(r2, size=10, bold=True, color=BLEU_FONCE)

    doc.add_paragraph("")

    # ========== MODE DE PAIEMENT ==========
    table_mode = doc.add_table(rows=1, cols=3)
    table_mode.autofit = True

    modes = [
        ("Mode de paiement", str(data.get("Mode_paiement", ""))),
        ("N° de chèque", str(data.get("Numero_cheque", "—") or "—")),
        ("Banque", str(data.get("Banque", ""))),
    ]
    for i, (label, val) in enumerate(modes):
        cell = table_mode.rows[0].cells[i]
        cell.text = ""
        shade_cell(cell, "0F2850")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p1 = cell.paragraphs[0]
        r1 = p1.add_run(label + "\n")
        set_run_font(r1, size=7, color=RGBColor(180, 200, 230))
        r2 = p1.add_run(val)
        set_run_font(r2, size=10, bold=True, color=RGBColor(255, 255, 255))
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")
    doc.add_paragraph("")

    # ========== SIGNATURES ==========
    table_sign = doc.add_table(rows=3, cols=3)
    table_sign.autofit = True

    labels_sign = ["Manager Logistics", "Comptabilité", "Bénéficiaire"]
    for i, lab in enumerate(labels_sign):
        cell = table_sign.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lab)
        set_run_font(r, size=9, bold=True, color=BLEU_FONCE)

        cell2 = table_sign.rows[1].cells[i]
        cell2.text = ""
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("\n\n________________")
        set_run_font(r2, size=10, color=GRIS_CLAIR)

        cell3 = table_sign.rows[2].cells[i]
        cell3.text = ""
        p3 = cell3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run("visa")
        set_run_font(r3, size=8, color=GRIS_CLAIR)

    doc.add_paragraph("")

    # ========== NOTE (rouge gras) ==========
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rn = note.add_run(
        "NB : Toute remise de chèque doit être confirmée par signature du bénéficiaire "
        "et accompagnée d'une copie du chèque et de la pièce d'identité pour fins de traçabilité."
    )
    set_run_font(rn, size=8, bold=True, color=RGBColor(180, 0, 0))

    # Sauvegarde
    no_trans = str(data.get("No_transaction", "PAIEMENT")).replace("/", "-").replace("\\", "-")
    nom_fichier = f"Bon_Paiement_{no_trans}.docx"
    chemin = os.path.join(DOSSIER_SORTIE, nom_fichier)
    doc.save(chemin)
    return chemin