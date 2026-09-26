import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import os
import re
from datetime import datetime
import shutil

st.set_page_config(page_title="Migration BC", page_icon="🔄", layout="wide")

if not st.session_state.get("authentication_status"):
    st.warning("Veuillez vous connecter depuis la page d'accueil.")
    st.stop()

username = st.session_state.get("username")
if username != "admin":
    st.error("⛔ Accès réservé à l'administrateur.")
    st.stop()

st.title("🔄 Migration des N° BC (ancien → nouveau format)")
st.markdown("---")

FICHIER_TRACKER = "SHEA_PURCHASE TRACKER.xlsx"
SUFFIXE = "-BC-SHEA/NAP-2026"

def convertir_bc(ancien):
    """BC-001 → 0001-BC-SHEA/NAP-2026 | BC-033 → 0033-BC-SHEA/NAP-2026"""
    if not ancien or pd.isna(ancien):
        return ancien
    s = str(ancien).strip().upper()
    match = re.search(r'BC-?(\d+)', s)
    if match:
        num = int(match.group(1))
        return f"{num:04d}{SUFFIXE}"
    # Déjà au nouveau format ?
    if re.match(r'^\d{3,}-BC-SHEA', s):
        return s
    return s

st.info("""
**Ce que fait cet outil :**
1. Crée une sauvegarde du tracker
2. Convertit tous les `BC-001`, `BC-002`… en `0001-BC-SHEA/NAP-2026`, `0002-…`
3. Met à jour les feuilles **Bons de Commande** et **Commandes Dechargees**
""")

if not os.path.exists(FICHIER_TRACKER):
    st.error(f"Fichier introuvable : {FICHIER_TRACKER}")
    st.stop()

# Aperçu
try:
    df_preview = pd.read_excel(FICHIER_TRACKER, sheet_name="Bons de Commande", header=3)
    df_preview.columns = [str(c).strip() for c in df_preview.columns]
    col_bc = next((c for c in df_preview.columns if "bc" in c.lower()), None)
    
    if col_bc:
        st.subheader("Aperçu de la conversion")
        echantillon = df_preview[[col_bc]].dropna().head(10).copy()
        echantillon["Nouveau format"] = echantillon[col_bc].apply(convertir_bc)
        st.dataframe(echantillon, width="stretch")
    else:
        st.warning("Colonne N° BC non trouvée dans Bons de Commande.")
except Exception as e:
    st.error(f"Erreur lecture : {e}")

st.markdown("---")

if st.button("🚀 Lancer la migration", type="primary"):
    with st.spinner("Migration en cours…"):
        try:
            # Sauvegarde
            backup_name = f"SHEA_PURCHASE_TRACKER_BACKUP_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            shutil.copy2(FICHIER_TRACKER, backup_name)
            st.success(f"Sauvegarde créée : `{backup_name}`")

            wb = load_workbook(FICHIER_TRACKER)
            
            # --- Feuille Bons de Commande ---
            if "Bons de Commande" in wb.sheetnames:
                ws = wb["Bons de Commande"]
                # Trouver la ligne d'en-tête (souvent ligne 4)
                header_row = 4
                headers = {cell.value: cell.column for cell in ws[header_row] if cell.value}
                col_bc_idx = None
                for h, idx in headers.items():
                    if h and "bc" in str(h).lower():
                        col_bc_idx = idx
                        break
                
                if col_bc_idx:
                    count = 0
                    for row in range(header_row + 1, ws.max_row + 1):
                        val = ws.cell(row=row, column=col_bc_idx).value
                        if val:
                            nouveau = convertir_bc(val)
                            if nouveau != val:
                                ws.cell(row=row, column=col_bc_idx).value = nouveau
                                count += 1
                    st.write(f"✅ Bons de Commande : {count} numéros convertis")
                else:
                    st.warning("Colonne N° BC non trouvée dans Bons de Commande")

            # --- Feuille Commandes Dechargees ---
            if "Commandes Dechargees" in wb.sheetnames:
                ws = wb["Commandes Dechargees"]
                header_row = 4
                headers = {cell.value: cell.column for cell in ws[header_row] if cell.value}
                col_bc_idx = None
                for h, idx in headers.items():
                    if h and ("bc" in str(h).lower() or "n° bc" in str(h).lower()):
                        col_bc_idx = idx
                        break
                
                if col_bc_idx:
                    count = 0
                    for row in range(header_row + 1, ws.max_row + 1):
                        val = ws.cell(row=row, column=col_bc_idx).value
                        if val:
                            nouveau = convertir_bc(val)
                            if nouveau != val:
                                ws.cell(row=row, column=col_bc_idx).value = nouveau
                                count += 1
                    st.write(f"✅ Commandes Dechargees : {count} numéros convertis")
                else:
                    st.warning("Colonne N° BC non trouvée dans Commandes Dechargees")

            wb.save(FICHIER_TRACKER)
            st.success("🎉 Migration terminée avec succès !")
            st.balloons()
            
        except Exception as e:
            st.error(f"Erreur pendant la migration : {e}")