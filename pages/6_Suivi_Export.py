import streamlit as st
import pandas as pd
from datetime import datetime
import os
from openpyxl import Workbook, load_workbook

st.set_page_config(page_title="Suivi Export", page_icon="🚢", layout="wide")

if not st.session_state.get("authentication_status"):
    st.warning("Veuillez vous connecter depuis la page d'accueil.")
    st.stop()

name = st.session_state.get("name")
username = st.session_state.get("username")

FICHIER = "exports_suivi.xlsx"

# ==================== INIT ====================
def init_fichier():
    if os.path.exists(FICHIER):
        return
    wb = Workbook()

    ws = wb.active
    ws.title = "Dossiers"
    ws.append([
        "N_Dossier", "Booking", "BL", "Produit", "Destination", "Transitaire",
        "Ligne_Maritime", "Statut", "Pct_Avancement", "Date_Creation",
        "Responsable", "Commentaire"
    ])

    # Modèle des pièces (personnalisable)
    ws2 = wb.create_sheet("Modele_Pieces")
    ws2.append(["Code", "Libelle", "Obligatoire", "Actif", "Ordre"])
    modeles_pieces = [
        ("P01", "Facture commerciale", "Oui", "Oui", 1),
        ("P02", "Packing list", "Oui", "Oui", 2),
        ("P03", "Certificat d'origine", "Oui", "Oui", 3),
        ("P04", "Certificat phytosanitaire", "Oui", "Oui", 4),
        ("P05", "COQ / COW", "Oui", "Oui", 5),
        ("P06", "Analyse labo / qualité", "Non", "Oui", 6),
        ("P07", "Booking confirmé", "Oui", "Oui", 7),
        ("P08", "VGM / note de poids", "Oui", "Oui", 8),
        ("P09", "Numéros de scellés", "Oui", "Oui", 9),
        ("P10", "Déclaration export (EX1)", "Oui", "Oui", 10),
        ("P11", "BAE / validation douane", "Oui", "Oui", 11),
        ("P12", "Manifeste export visé", "Oui", "Oui", 12),
        ("P13", "B/L draft", "Oui", "Oui", 13),
        ("P14", "B/L final / Telex", "Oui", "Oui", 14),
        ("P15", "BESC / CTN (si requis)", "Non", "Oui", 15),
    ]
    for m in modeles_pieces:
        ws2.append(list(m))

    # Modèle des étapes (personnalisable)
    ws3 = wb.create_sheet("Modele_Etapes")
    ws3.append(["N_Etape", "Libelle", "Actif", "Ordre"])
    modeles_etapes = [
        (1, "Contrat / confirmation vente", "Oui", 1),
        (2, "Choix transitaire + mandat", "Oui", 2),
        (3, "Booking obtenu", "Oui", 3),
        (4, "Conteneurs positionnés", "Oui", 4),
        (5, "Empotage terminé", "Oui", 5),
        (6, "Pesage + VGM", "Oui", 6),
        (7, "Scellés posés", "Oui", 7),
        (8, "Pièces complètes (checklist)", "Oui", 8),
        (9, "Déclaration douane déposée", "Oui", 9),
        (10, "Contrôle douane OK / BAE", "Oui", 10),
        (11, "Transfert au terminal", "Oui", 11),
        (12, "Manifeste export visé", "Oui", 12),
        (13, "Embarquement (vu à bord)", "Oui", 13),
        (14, "B/L final / Telex reçu", "Oui", 14),
        (15, "Envoi docs à acheteur / banque", "Oui", 15),
        (16, "Clôture dossier", "Oui", 16),
    ]
    for m in modeles_etapes:
        ws3.append(list(m))

    ws4 = wb.create_sheet("Pieces")
    ws4.append([
        "N_Dossier", "Code_Piece", "Libelle", "Obligatoire", "Statut",
        "Date_Validation", "Valide_Par", "Commentaire"
    ])

    ws5 = wb.create_sheet("Etapes")
    ws5.append([
        "N_Dossier", "N_Etape", "Libelle", "Statut", "Date_Debut",
        "Date_Fin", "Commentaire"
    ])

    ws6 = wb.create_sheet("Historique")
    ws6.append(["Date_Heure", "N_Dossier", "Utilisateur", "Action", "Detail"])

    wb.save(FICHIER)

def charger_sheet(nom):
    init_fichier()
    try:
        return pd.read_excel(FICHIER, sheet_name=nom)
    except Exception:
        return pd.DataFrame()

def charger_modele_pieces():
    df = charger_sheet("Modele_Pieces")
    if df.empty:
        return df
    df = df[df["Actif"].astype(str).str.upper() == "OUI"]
    if "Ordre" in df.columns:
        df = df.sort_values("Ordre")
    return df

def charger_modele_etapes():
    df = charger_sheet("Modele_Etapes")
    if df.empty:
        return df
    df = df[df["Actif"].astype(str).str.upper() == "OUI"]
    if "Ordre" in df.columns:
        df = df.sort_values("Ordre")
    return df

def charger_dossiers():
    return charger_sheet("Dossiers")

def charger_pieces(n_dossier=None):
    df = charger_sheet("Pieces")
    if n_dossier and not df.empty:
        df = df[df["N_Dossier"].astype(str) == str(n_dossier)]
    return df

def charger_etapes(n_dossier=None):
    df = charger_sheet("Etapes")
    if n_dossier and not df.empty:
        df = df[df["N_Dossier"].astype(str) == str(n_dossier)]
    return df

def prochain_numero():
    df = charger_dossiers()
    if df.empty:
        return "EXP-2026-0001"
    nums = []
    for v in df["N_Dossier"].dropna().astype(str):
        if "EXP-2026-" in v:
            try:
                nums.append(int(v.split("-")[-1]))
            except Exception:
                pass
    n = max(nums) + 1 if nums else 1
    return f"EXP-2026-{n:04d}"

def ajouter_historique(n_dossier, action, detail):
    wb = load_workbook(FICHIER)
    ws = wb["Historique"]
    ws.append([
        datetime.now().strftime("%d/%m/%Y %H:%M"),
        n_dossier, name, action, detail
    ])
    wb.save(FICHIER)

def creer_dossier(data):
    n_dossier = prochain_numero()
    wb = load_workbook(FICHIER)

    ws = wb["Dossiers"]
    ws.append([
        n_dossier,
        data.get("Booking", ""),
        data.get("BL", ""),
        data.get("Produit", ""),
        data.get("Destination", ""),
        data.get("Transitaire", ""),
        data.get("Ligne_Maritime", ""),
        "EN COURS",
        0,
        datetime.now().strftime("%d/%m/%Y"),
        name,
        data.get("Commentaire", "")
    ])

    # Pièces selon le modèle actuel (personnalisé)
    modele_p = charger_modele_pieces()
    ws2 = wb["Pieces"]
    for _, p in modele_p.iterrows():
        ws2.append([
            n_dossier,
            str(p.get("Code", "")),
            str(p.get("Libelle", "")),
            str(p.get("Obligatoire", "Oui")),
            "MANQUANT", "", "", ""
        ])

    # Étapes selon le modèle actuel
    modele_e = charger_modele_etapes()
    ws3 = wb["Etapes"]
    for _, e in modele_e.iterrows():
        ws3.append([
            n_dossier,
            int(e.get("N_Etape", 0)),
            str(e.get("Libelle", "")),
            "À FAIRE", "", "", ""
        ])

    wb.save(FICHIER)
    ajouter_historique(n_dossier, "Création", f"Dossier créé – Booking {data.get('Booking', '')}")
    return n_dossier

def maj_piece(n_dossier, code_piece, statut, commentaire=""):
    wb = load_workbook(FICHIER)
    ws = wb["Pieces"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(n_dossier) and str(row[1].value) == str(code_piece):
            row[4].value = statut
            if statut == "OK":
                row[5].value = datetime.now().strftime("%d/%m/%Y")
                row[6].value = name
            row[7].value = commentaire
            break
    wb.save(FICHIER)
    ajouter_historique(n_dossier, "Pièce", f"{code_piece} → {statut}")
    recalculer_avancement(n_dossier)

def maj_etape(n_dossier, n_etape, statut, commentaire=""):
    wb = load_workbook(FICHIER)
    ws = wb["Etapes"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(n_dossier) and int(row[1].value) == int(n_etape):
            row[3].value = statut
            if statut == "EN COURS" and not row[4].value:
                row[4].value = datetime.now().strftime("%d/%m/%Y")
            if statut == "FAIT":
                row[5].value = datetime.now().strftime("%d/%m/%Y")
            row[6].value = commentaire
            break
    wb.save(FICHIER)
    ajouter_historique(n_dossier, "Étape", f"Étape {n_etape} → {statut}")
    recalculer_avancement(n_dossier)

def recalculer_avancement(n_dossier):
    df_e = charger_etapes(n_dossier)
    if df_e.empty:
        return
    total = len(df_e)
    faits = (df_e["Statut"].astype(str) == "FAIT").sum()
    pct = int(round(100 * faits / total)) if total else 0

    if (df_e["Statut"].astype(str) == "BLOQUÉ").any():
        statut = "BLOQUÉ"
    elif pct >= 100:
        statut = "CLOTURÉ"
    else:
        # Embarquement = étape dont le libellé contient "Embarquement"
        emb = df_e[df_e["Libelle"].astype(str).str.contains("Embarquement", case=False, na=False)]
        if not emb.empty and (emb["Statut"].astype(str) == "FAIT").any():
            statut = "EMBARQUÉ"
        else:
            statut = "EN COURS"

    wb = load_workbook(FICHIER)
    ws = wb["Dossiers"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(n_dossier):
            row[7].value = statut
            row[8].value = pct
            break
    wb.save(FICHIER)

def pieces_manquantes_obligatoires(n_dossier):
    df = charger_pieces(n_dossier)
    if df.empty:
        return []
    mask = (
        (df["Obligatoire"].astype(str).str.upper() == "OUI") &
        (df["Statut"].astype(str).str.upper() != "OK")
    )
    return df.loc[mask, "Libelle"].tolist()

def etapes_bloquees(n_dossier):
    df = charger_etapes(n_dossier)
    if df.empty:
        return []
    return df[df["Statut"].astype(str).str.upper() == "BLOQUÉ"][
        ["N_Etape", "Libelle", "Commentaire"]
    ].to_dict("records")

def ajouter_modele_piece(code, libelle, obligatoire, ordre):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Pieces"]
    ws.append([code, libelle, "Oui" if obligatoire else "Non", "Oui", ordre])
    wb.save(FICHIER)

def ajouter_modele_etape(n_etape, libelle, ordre):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Etapes"]
    ws.append([n_etape, libelle, "Oui", ordre])
    wb.save(FICHIER)

def desactiver_modele_piece(code):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Pieces"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(code):
            row[3].value = "Non"  # Actif = Non
            break
    wb.save(FICHIER)

def desactiver_modele_etape(n_etape):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Etapes"]
    for row in ws.iter_rows(min_row=2):
        if int(row[0].value) == int(n_etape):
            row[2].value = "Non"
            break
    wb.save(FICHIER)

def renommer_modele_piece(code, nouveau_libelle, obligatoire):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Pieces"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(code):
            row[1].value = nouveau_libelle
            row[2].value = "Oui" if obligatoire else "Non"
            break
    wb.save(FICHIER)

def renommer_modele_etape(n_etape, nouveau_libelle):
    wb = load_workbook(FICHIER)
    ws = wb["Modele_Etapes"]
    for row in ws.iter_rows(min_row=2):
        if int(row[0].value) == int(n_etape):
            row[1].value = nouveau_libelle
            break
    wb.save(FICHIER)

# ==================== UI ====================
st.title("🚢 Suivi Export – par B/L / Booking")
st.caption("Processus adaptable à votre réalité • Pièces et étapes personnalisables")
st.markdown("---")

df_dossiers = charger_dossiers()
nb_cours = nb_bloque = nb_embarque = nb_cloture = 0
if not df_dossiers.empty and "Statut" in df_dossiers.columns:
    nb_cours = (df_dossiers["Statut"].astype(str) == "EN COURS").sum()
    nb_bloque = (df_dossiers["Statut"].astype(str) == "BLOQUÉ").sum()
    nb_embarque = (df_dossiers["Statut"].astype(str) == "EMBARQUÉ").sum()
    nb_cloture = (df_dossiers["Statut"].astype(str) == "CLOTURÉ").sum()

c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Total dossiers", len(df_dossiers))
c2.metric("En cours", nb_cours)
c3.metric("Bloqués", nb_bloque)
c4.metric("Embarqués", nb_embarque)
c5.metric("Clôturés", nb_cloture)

st.markdown("---")

onglet1, onglet2, onglet3, onglet4 = st.tabs([
    "📋 Liste des dossiers",
    "➕ Nouveau dossier",
    "🔍 Fiche détaillée",
    "⚙️ Configurer pièces & étapes"
])

# ========== LISTE ==========
with onglet1:
    st.subheader("Tous les dossiers export")
    if df_dossiers.empty:
        st.info("Aucun dossier. Créez-en un dans « Nouveau dossier ».")
    else:
        filtre = st.selectbox("Filtrer par statut", ["Tous", "EN COURS", "BLOQUÉ", "EMBARQUÉ", "CLOTURÉ"])
        df_aff = df_dossiers.copy()
        if filtre != "Tous":
            df_aff = df_aff[df_aff["Statut"].astype(str) == filtre]
        st.dataframe(df_aff, width="stretch", height=350)

        bloques = df_dossiers[df_dossiers["Statut"].astype(str) == "BLOQUÉ"]
        if not bloques.empty:
            st.error(f"⚠️ {len(bloques)} dossier(s) BLOQUÉ(S)")
            for _, r in bloques.iterrows():
                st.write(f"• **{r['N_Dossier']}** | BL: {r.get('BL', '—')} | {r.get('Transitaire', '—')}")

# ========== NOUVEAU ==========
with onglet2:
    st.subheader("Créer un nouveau dossier export")
    st.caption("Le dossier reprend **vos** pièces et étapes actives (onglet Configurer).")

    if st.session_state.get("export_success"):
        st.success(st.session_state["export_success"])
        del st.session_state["export_success"]

    with st.form("form_nouveau", clear_on_submit=True):
        col1, col2 = st.columns(2)
        with col1:
            booking = st.text_input("N° Booking *")
            bl = st.text_input("N° B/L (si déjà connu)")
            produit = st.selectbox("Produit", ["SHEA", "SOYABEAN", "AUTRE"])
            destination = st.text_input("Destination *", placeholder="Ex: Chine, Europe…")
        with col2:
            transitaire = st.text_input("Transitaire *", placeholder="Ex: TRANS SOHA")
            ligne = st.text_input("Ligne maritime", placeholder="Ex: MSC, Maersk")
            commentaire = st.text_area("Commentaire", height=80)

        submitted = st.form_submit_button("Créer le dossier", type="primary")
        if submitted:
            if not booking.strip() or not destination.strip() or not transitaire.strip():
                st.error("Booking, Destination et Transitaire sont obligatoires.")
            else:
                n = creer_dossier({
                    "Booking": booking.strip(),
                    "BL": bl.strip(),
                    "Produit": produit,
                    "Destination": destination.strip(),
                    "Transitaire": transitaire.strip(),
                    "Ligne_Maritime": ligne.strip(),
                    "Commentaire": commentaire.strip()
                })
                st.session_state["export_success"] = (
                    f"✅ Dossier **{n}** créé avec succès "
                    f"(Booking : {booking.strip()}). "
                    f"Checklist selon votre configuration."
                )
                st.rerun()

# ========== FICHE ==========
with onglet3:
    st.subheader("Fiche détaillée d’un dossier")
    if df_dossiers.empty:
        st.info("Aucun dossier à afficher.")
    else:
        options = df_dossiers["N_Dossier"].astype(str).tolist()
        labels = []
        for _, r in df_dossiers.iterrows():
            labels.append(f"{r['N_Dossier']} | BL: {r.get('BL', '') or '—'} | {r.get('Statut', '')}")
        choix_idx = st.selectbox("Choisir le dossier", range(len(options)), format_func=lambda i: labels[i])
        n_dossier = options[choix_idx]
        row = df_dossiers[df_dossiers["N_Dossier"].astype(str) == n_dossier].iloc[0]

        st.markdown(f"### {n_dossier} — **{row.get('Statut', '')}** ({row.get('Pct_Avancement', 0)} %)")
        col_a, col_b, col_c = st.columns(3)
        col_a.write(f"**Booking :** {row.get('Booking', '—')}")
        col_a.write(f"**B/L :** {row.get('BL', '—')}")
        col_b.write(f"**Produit :** {row.get('Produit', '—')}")
        col_b.write(f"**Destination :** {row.get('Destination', '—')}")
        col_c.write(f"**Transitaire :** {row.get('Transitaire', '—')}")
        col_c.write(f"**Ligne :** {row.get('Ligne_Maritime', '—')}")

        manquantes = pieces_manquantes_obligatoires(n_dossier)
        bloques_e = etapes_bloquees(n_dossier)
        if manquantes:
            st.warning("⚠️ **Pièces obligatoires manquantes :** " + " · ".join(manquantes))
        if bloques_e:
            for b in bloques_e:
                st.error(f"🔴 **Étape {b['N_Etape']} bloquée :** {b['Libelle']} — {b.get('Commentaire', '')}")

        st.markdown("---")
        st.markdown("#### 📎 Checklist des pièces")
        df_p = charger_pieces(n_dossier)
        for _, p in df_p.iterrows():
            code = p["Code_Piece"]
            lib = p["Libelle"]
            oblig = p["Obligatoire"]
            statut = str(p["Statut"])
            colx, coly, colz = st.columns([4, 2, 2])
            with colx:
                badge = "🔴" if statut != "OK" and str(oblig).upper() == "OUI" else ("✅" if statut == "OK" else "⚪")
                st.write(f"{badge} **{lib}**" + (" *(obligatoire)*" if str(oblig).upper() == "OUI" else ""))
            with coly:
                nouveau = st.selectbox(
                    "Statut", ["MANQUANT", "OK", "N/A"],
                    index=["MANQUANT", "OK", "N/A"].index(statut) if statut in ["MANQUANT", "OK", "N/A"] else 0,
                    key=f"piece_{n_dossier}_{code}"
                )
            with colz:
                if st.button("Enregistrer", key=f"btn_p_{n_dossier}_{code}"):
                    maj_piece(n_dossier, code, nouveau)
                    st.rerun()

        st.markdown("---")
        st.markdown("#### 📍 Étapes du process")
        df_e = charger_etapes(n_dossier)
        for _, e in df_e.iterrows():
            n_et = int(e["N_Etape"])
            lib = e["Libelle"]
            statut = str(e["Statut"])
            icon = {"À FAIRE": "⬜", "EN COURS": "🔄", "FAIT": "✅", "BLOQUÉ": "🔴", "N/A": "➖"}.get(statut, "⬜")
            col1, col2, col3 = st.columns([4, 2, 3])
            with col1:
                st.write(f"{icon} **{n_et}. {lib}**")
                if e.get("Date_Debut") or e.get("Date_Fin"):
                    st.caption(f"Début: {e.get('Date_Debut', '—')} | Fin: {e.get('Date_Fin', '—')}")
            with col2:
                nouveau_e = st.selectbox(
                    "Statut", ["À FAIRE", "EN COURS", "FAIT", "BLOQUÉ", "N/A"],
                    index=["À FAIRE", "EN COURS", "FAIT", "BLOQUÉ", "N/A"].index(statut)
                    if statut in ["À FAIRE", "EN COURS", "FAIT", "BLOQUÉ", "N/A"] else 0,
                    key=f"etape_{n_dossier}_{n_et}"
                )
            with col3:
                com = st.text_input("Commentaire", value=str(e.get("Commentaire", "") or ""), key=f"com_{n_dossier}_{n_et}")
                if st.button("Maj étape", key=f"btn_e_{n_dossier}_{n_et}"):
                    # Si libellé contient "pièces complètes" → vérifier obligatoires
                    if "pièce" in lib.lower() and nouveau_e == "FAIT":
                        manq = pieces_manquantes_obligatoires(n_dossier)
                        if manq:
                            st.error("Pièces obligatoires manquantes : " + ", ".join(manq))
                        else:
                            maj_etape(n_dossier, n_et, nouveau_e, com)
                            st.rerun()
                    else:
                        maj_etape(n_dossier, n_et, nouveau_e, com)
                        st.rerun()

        st.markdown("---")
        st.markdown("#### ✏️ Mettre à jour le N° B/L")
        nouveau_bl = st.text_input("N° B/L", value=str(row.get("BL", "") or ""))
        if st.button("Enregistrer le B/L"):
            wb = load_workbook(FICHIER)
            ws = wb["Dossiers"]
            for r in ws.iter_rows(min_row=2):
                if str(r[0].value) == str(n_dossier):
                    r[2].value = nouveau_bl.strip()
                    break
            wb.save(FICHIER)
            ajouter_historique(n_dossier, "BL", f"B/L mis à jour : {nouveau_bl}")
            st.success("B/L enregistré.")
            st.rerun()

# ========== CONFIG ==========
with onglet4:
    st.subheader("⚙️ Configurer le process selon votre réalité")
    st.info(
        "Les **nouveaux dossiers** utiliseront cette configuration. "
        "Les dossiers déjà créés gardent leur checklist d’origine."
    )

    conf1, conf2 = st.tabs(["📎 Pièces (documents)", "📍 Étapes (process)"])

    # --- Pièces ---
    with conf1:
        st.markdown("##### Pièces actives")
        df_mp = charger_sheet("Modele_Pieces")
        if not df_mp.empty:
            actifs = df_mp[df_mp["Actif"].astype(str).str.upper() == "OUI"]
            st.dataframe(actifs, width="stretch", height=250)

        st.markdown("##### Ajouter une pièce")
        with st.form("add_piece", clear_on_submit=True):
            c1, c2, c3 = st.columns(3)
            with c1:
                code = st.text_input("Code *", placeholder="P16")
            with c2:
                lib = st.text_input("Nom de la pièce *", placeholder="Ex: Certificat fumigation")
            with c3:
                ordre = st.number_input("Ordre", min_value=1, value=20)
            oblig = st.checkbox("Obligatoire", value=True)
            if st.form_submit_button("➕ Ajouter la pièce"):
                if code.strip() and lib.strip():
                    ajouter_modele_piece(code.strip(), lib.strip(), oblig, int(ordre))
                    st.success(f"Pièce **{lib}** ajoutée.")
                    st.rerun()
                else:
                    st.error("Code et nom obligatoires.")

        st.markdown("##### Modifier / désactiver une pièce")
        if not df_mp.empty:
            codes = df_mp[df_mp["Actif"].astype(str).str.upper() == "OUI"]["Code"].astype(str).tolist()
            if codes:
                code_sel = st.selectbox("Pièce", codes, key="mod_piece")
                row_p = df_mp[df_mp["Code"].astype(str) == code_sel].iloc[0]
                new_lib = st.text_input("Nouveau nom", value=str(row_p["Libelle"]), key="ren_piece")
                new_ob = st.checkbox("Obligatoire", value=str(row_p["Obligatoire"]).upper() == "OUI", key="ob_piece")
                col_a, col_b = st.columns(2)
                with col_a:
                    if st.button("💾 Enregistrer le nom"):
                        renommer_modele_piece(code_sel, new_lib.strip(), new_ob)
                        st.success("Modifié.")
                        st.rerun()
                with col_b:
                    if st.button("🗑️ Désactiver cette pièce"):
                        desactiver_modele_piece(code_sel)
                        st.success("Pièce désactivée (plus utilisée pour les nouveaux dossiers).")
                        st.rerun()

    # --- Étapes ---
    with conf2:
        st.markdown("##### Étapes actives")
        df_me = charger_sheet("Modele_Etapes")
        if not df_me.empty:
            actifs_e = df_me[df_me["Actif"].astype(str).str.upper() == "OUI"]
            st.dataframe(actifs_e, width="stretch", height=250)

        st.markdown("##### Ajouter une étape")
        with st.form("add_etape", clear_on_submit=True):
            c1, c2 = st.columns(2)
            with c1:
                # Prochain numéro
                max_n = int(df_me["N_Etape"].max()) + 1 if not df_me.empty else 1
                n_et = st.number_input("N° étape", min_value=1, value=max_n)
            with c2:
                ordre_e = st.number_input("Ordre d'affichage", min_value=1, value=max_n)
            lib_e = st.text_input("Nom de l'étape *", placeholder="Ex: Inspection phytosanitaire port")
            if st.form_submit_button("➕ Ajouter l'étape"):
                if lib_e.strip():
                    ajouter_modele_etape(int(n_et), lib_e.strip(), int(ordre_e))
                    st.success(f"Étape **{lib_e}** ajoutée.")
                    st.rerun()
                else:
                    st.error("Nom obligatoire.")

        st.markdown("##### Modifier / désactiver une étape")
        if not df_me.empty:
            actifs_e = df_me[df_me["Actif"].astype(str).str.upper() == "OUI"]
            if not actifs_e.empty:
                opts = [f"{int(r['N_Etape'])} – {r['Libelle']}" for _, r in actifs_e.iterrows()]
                sel = st.selectbox("Étape", opts, key="mod_etape")
                n_sel = int(sel.split("–")[0].strip())
                row_e = df_me[df_me["N_Etape"] == n_sel].iloc[0]
                new_lib_e = st.text_input("Nouveau nom", value=str(row_e["Libelle"]), key="ren_etape")
                col_a, col_b = st.columns(2)
                with col_a:
                    if st.button("💾 Enregistrer le nom", key="save_etape"):
                        renommer_modele_etape(n_sel, new_lib_e.strip())
                        st.success("Modifié.")
                        st.rerun()
                with col_b:
                    if st.button("🗑️ Désactiver cette étape", key="del_etape"):
                        desactiver_modele_etape(n_sel)
                        st.success("Étape désactivée.")
                        st.rerun()

# Téléchargement
st.markdown("---")
if os.path.exists(FICHIER):
    with open(FICHIER, "rb") as f:
        st.download_button(
            "⬇️ Télécharger exports_suivi.xlsx",
            data=f,
            file_name="exports_suivi.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

st.caption("NAP SARL – Suivi export adaptable • Config pièces & étapes selon votre process")