import streamlit as st
import pandas as pd
from datetime import datetime
import os
import re
from openpyxl import load_workbook

st.set_page_config(page_title="Déchargements", page_icon="🚛", layout="wide")

if not st.session_state.get("authentication_status"):
    st.warning("Veuillez vous connecter depuis la page d'accueil.")
    st.stop()

name = st.session_state.get("name")
username = st.session_state.get("username")

# ==================== CONFIGURATION ====================
FICHIER_BONS = "donnees_bons.xlsx"
FICHIER_TRACKER = "SHEA_PURCHASE TRACKER.xlsx"
FICHIER_DECHARGEMENTS = "dechargements.xlsx"
FICHIER_FRS = "fournisseurs.xlsx"
# =======================================================

st.title("🚛 Enregistrement des Déchargements (Tickets de Pont)")
st.caption("Montant dû = Poids net (kg) × PU fournisseur (FCFA/kg)")
st.markdown("---")

def charger_bons_app():
    if not os.path.exists(FICHIER_BONS):
        return pd.DataFrame()
    return pd.read_excel(FICHIER_BONS)

def charger_bons_tracker():
    if not os.path.exists(FICHIER_TRACKER):
        return pd.DataFrame()
    try:
        df = pd.read_excel(FICHIER_TRACKER, sheet_name="Bons de Commande", header=3)
        df.columns = [str(c).strip() for c in df.columns]
        return df
    except Exception:
        return pd.DataFrame()

def charger_dechargements():
    if os.path.exists(FICHIER_DECHARGEMENTS):
        return pd.read_excel(FICHIER_DECHARGEMENTS)
    return pd.DataFrame(columns=[
        "Date_Dechargement", "N_BC", "Code_Frs", "Fournisseur", "N_Camion",
        "Weighment_Ticket", "Poids_Net_KG", "Nb_Sacs", "PU_FCFA_KG",
        "Montant_Du", "Montant_Paye", "Solde_Restant", "Statut_Paiement",
        "Bon_Paiement", "Statut_Livraison", "Lieu_Stockage", "Utilisateur"
    ])

def charger_pu_fournisseur(code_frs=None, nom_frs=None):
    """Retourne PU_FCFA_KG depuis fournisseurs.xlsx (par code ou nom)."""
    if not os.path.exists(FICHIER_FRS):
        return 0.0
    try:
        df = pd.read_excel(FICHIER_FRS)
    except Exception:
        return 0.0
    if df.empty or "PU_FCFA_KG" not in df.columns:
        return 0.0

    row = None
    if code_frs:
        m = df[df["Code"].astype(str).str.strip().str.upper() == str(code_frs).strip().upper()]
        if not m.empty:
            row = m.iloc[0]
    if row is None and nom_frs:
        m = df[df["Nom_Fournisseur"].astype(str).str.strip().str.upper() == str(nom_frs).strip().upper()]
        if not m.empty:
            row = m.iloc[0]
    if row is None:
        return 0.0
    try:
        return float(row.get("PU_FCFA_KG") or 0)
    except Exception:
        return 0.0

def normaliser_camion(val):
    if not val or (isinstance(val, float) and pd.isna(val)):
        return ""
    s = str(val).upper().strip()
    s = re.sub(r'[\s\-/\\]+', '', s)
    return s

def extraire_masse_mt(val):
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).upper().replace(",", ".")
    m = re.search(r'(\d+(?:\.\d+)?)\s*MT', s)
    if m:
        return float(m.group(1))
    m = re.search(r'(\d+(?:\.\d+)?)\s*KG', s)
    if m:
        return float(m.group(1)) / 1000.0
    m = re.search(r'(\d+(?:\.\d+)?)', s)
    if m:
        n = float(m.group(1))
        if n > 200:
            return n / 1000.0
        return n
    return None

def match_camion(camion_saisi_norm, camion_row_raw):
    camion_row = normaliser_camion(camion_row_raw)
    if not camion_row or not camion_saisi_norm:
        return False
    if camion_saisi_norm in camion_row or camion_row in camion_saisi_norm:
        return True
    parties = re.findall(r'[A-Z0-9]{5,}', camion_saisi_norm)
    return any(p in camion_row for p in parties)

def rechercher_par_camion(numero_camion, masse_saisie, df_app, df_tracker, df_dech):
    numero_camion = str(numero_camion).strip()
    if not numero_camion:
        return [], "Veuillez saisir un N° Camion."

    camion_norm = normaliser_camion(numero_camion)
    masse_mt = extraire_masse_mt(masse_saisie) if masse_saisie else None
    resultats = []
    alertes = []

    if not df_app.empty and "Numero_Tracteur" in df_app.columns:
        for _, row in df_app.iterrows():
            if not match_camion(camion_norm, row.get("Numero_Tracteur", "")):
                continue
            qte = row.get("Quantite", "")
            masse_bon = extraire_masse_mt(qte)
            score_masse = None
            if masse_mt is not None and masse_bon is not None:
                diff = abs(masse_mt - masse_bon)
                score_masse = "OK" if diff <= 2.0 else f"ÉCART {diff:.1f} MT"

            code_frs = str(row.get("Code_Fournisseur", "") or "").strip()
            fourn = str(row.get("Nom_Fournisseur", "") or "").strip()
            pu = charger_pu_fournisseur(code_frs, fourn)

            resultats.append({
                "source": "App (nouveau format)",
                "Reference": str(row.get("Reference", "")),
                "Code_Frs": code_frs,
                "Fournisseur": fourn,
                "Date": str(row.get("Date_Emission", "")),
                "Quantite": str(qte),
                "Masse_MT": masse_bon,
                "Camion": str(row.get("Numero_Tracteur", "")),
                "Chauffeur": str(row.get("Chauffeur", "")),
                "PU_FCFA_KG": pu,
                "Statut_paiement": "—",
                "score_masse": score_masse,
                "priorite": 1,
                "deja_decharge": False,
            })

    if not df_tracker.empty:
        col_camion = next((c for c in df_tracker.columns if "camion" in c.lower()), None)
        col_bc = next((c for c in df_tracker.columns if "bc" in c.lower()), None)
        col_fourn = next((c for c in df_tracker.columns if "fournisseur" in c.lower()), None)
        col_qte = next((c for c in df_tracker.columns if "qty" in c.lower() or "quantit" in c.lower()), None)
        col_statut = next((c for c in df_tracker.columns if "statut" in c.lower() and "livr" in c.lower()), None)
        col_code = next((c for c in df_tracker.columns if "code" in c.lower() and "frs" in c.lower()), None)

        if col_camion:
            for _, row in df_tracker.iterrows():
                if not match_camion(camion_norm, row.get(col_camion, "")):
                    continue
                qte = row.get(col_qte, "") if col_qte else ""
                masse_bon = extraire_masse_mt(qte)
                score_masse = None
                if masse_mt is not None and masse_bon is not None:
                    diff = abs(masse_mt - masse_bon)
                    score_masse = "OK" if diff <= 2.0 else f"ÉCART {diff:.1f} MT"

                statut_liv = str(row.get(col_statut, "")) if col_statut else ""
                priorite = 2 if "DECHARGE" in statut_liv.upper() else 1
                fourn = str(row.get(col_fourn, "")) if col_fourn else ""
                code_frs = str(row.get(col_code, "")) if col_code else ""
                pu = charger_pu_fournisseur(code_frs, fourn)

                resultats.append({
                    "source": "Tracker",
                    "Reference": str(row.get(col_bc, "")) if col_bc else "",
                    "Code_Frs": code_frs,
                    "Fournisseur": fourn,
                    "Date": str(row.get("Date", "")),
                    "Quantite": str(qte),
                    "Masse_MT": masse_bon,
                    "Camion": str(row.get(col_camion, "")),
                    "Chauffeur": str(row.get("Contact Chauffeur", "")),
                    "PU_FCFA_KG": pu,
                    "Statut_paiement": statut_liv,
                    "score_masse": score_masse,
                    "priorite": priorite,
                    "deja_decharge": "DECHARGE" in statut_liv.upper(),
                })

    if not df_dech.empty and "N_Camion" in df_dech.columns:
        for _, row in df_dech.iterrows():
            if not match_camion(camion_norm, row.get("N_Camion", "")):
                continue
            statut_p = str(row.get("Statut_Paiement", "")).upper()
            masse_d = extraire_masse_mt(row.get("Poids_Net_KG", ""))
            score_masse = None
            if masse_mt is not None and masse_d is not None:
                diff = abs(masse_mt - masse_d)
                score_masse = "OK" if diff <= 2.0 else f"ÉCART {diff:.1f} MT"

            priorite = 0 if statut_p in ("NON PAYE", "PARTIEL") else 3
            try:
                pu = float(row.get("PU_FCFA_KG") or 0)
            except Exception:
                pu = 0.0

            resultats.append({
                "source": "Déchargement déjà saisi",
                "Reference": str(row.get("N_BC", "")),
                "Code_Frs": str(row.get("Code_Frs", "") or ""),
                "Fournisseur": str(row.get("Fournisseur", "")),
                "Date": str(row.get("Date_Dechargement", "")),
                "Quantite": f"{row.get('Poids_Net_KG', '')} KG",
                "Masse_MT": masse_d,
                "Camion": str(row.get("N_Camion", "")),
                "Chauffeur": "",
                "PU_FCFA_KG": pu,
                "Statut_paiement": statut_p,
                "score_masse": score_masse,
                "priorite": priorite,
                "deja_decharge": True,
            })

    resultats.sort(key=lambda x: (x.get("priorite", 9), str(x.get("Reference", ""))))

    if not resultats:
        alertes.append("⚠️ Aucun bon trouvé pour ce N° Camion. Vérifiez le numéro ou créez d'abord le bon de commande.")
    else:
        refs = [r["Reference"] for r in resultats if r.get("Reference")]
        if len(set(refs)) > 1:
            alertes.append(
                f"ℹ️ Ce camion apparaît sur **{len(set(refs))} bons différents**. "
                "Vérifiez le fournisseur et la masse pour choisir le bon."
            )
        ecarts = [r for r in resultats if r.get("score_masse") and "ÉCART" in str(r.get("score_masse", ""))]
        if masse_mt is not None and ecarts:
            alertes.append(
                f"⚠️ Écart de masse détecté (saisie ≈ {masse_mt:.2f} MT). "
                "Contrôlez le ticket de pont avant d'enregistrer."
            )
        deja_payes = [r for r in resultats if r.get("deja_decharge") and r.get("Statut_paiement") == "PAYE"]
        if deja_payes and not any(r.get("priorite") == 1 for r in resultats):
            alertes.append(
                "⚠️ Ce camion a déjà des déchargements **PAYÉS**. "
                "Si c'est un nouveau passage, créez d'abord un nouveau bon de commande."
            )

    return resultats, " | ".join(alertes) if alertes else ""

def sauvegarder_dechargement(ligne):
    df = charger_dechargements()
    df = pd.concat([df, pd.DataFrame([ligne])], ignore_index=True)
    df.to_excel(FICHIER_DECHARGEMENTS, index=False)

    if not os.path.exists(FICHIER_TRACKER):
        return True, "Enregistré dans dechargements.xlsx (tracker introuvable)"

    try:
        wb = load_workbook(FICHIER_TRACKER)
        if "Commandes Dechargees" not in wb.sheetnames:
            return True, "Enregistré dans dechargements.xlsx (feuille Commandes Dechargees absente)"

        ws = wb["Commandes Dechargees"]
        next_row = ws.max_row + 1
        while next_row > 5 and all(
            ws.cell(row=next_row - 1, column=c).value is None for c in range(1, 18)
        ):
            next_row -= 1

        ws.cell(row=next_row, column=1).value = ligne.get("Date_Dechargement")
        ws.cell(row=next_row, column=2).value = ligne.get("N_BC")
        ws.cell(row=next_row, column=4).value = ligne.get("Code_Frs")
        ws.cell(row=next_row, column=5).value = ligne.get("Fournisseur")
        ws.cell(row=next_row, column=6).value = ligne.get("N_Camion")
        ws.cell(row=next_row, column=7).value = ligne.get("Weighment_Ticket")
        ws.cell(row=next_row, column=8).value = ligne.get("Poids_Net_KG")
        ws.cell(row=next_row, column=9).value = ligne.get("Nb_Sacs")
        ws.cell(row=next_row, column=10).value = ligne.get("PU_FCFA_KG")
        ws.cell(row=next_row, column=11).value = ligne.get("Montant_Du")
        ws.cell(row=next_row, column=12).value = ligne.get("Montant_Paye")
        ws.cell(row=next_row, column=13).value = ligne.get("Solde_Restant")
        ws.cell(row=next_row, column=14).value = ligne.get("Statut_Paiement")
        ws.cell(row=next_row, column=15).value = ligne.get("Bon_Paiement")
        ws.cell(row=next_row, column=16).value = ligne.get("Statut_Livraison")
        ws.cell(row=next_row, column=17).value = ligne.get("Lieu_Stockage")

        wb.save(FICHIER_TRACKER)
        return True, "Enregistré dans dechargements.xlsx **et** dans le tracker"
    except Exception as e:
        return True, f"Enregistré dans dechargements.xlsx (erreur tracker : {e})"

# ---------- Chargement ----------
df_app = charger_bons_app()
df_tracker = charger_bons_tracker()
df_dech = charger_dechargements()

# ---------- Étape 1 ----------
st.subheader("1️⃣ Vérification Camion + Masse")

col_r1, col_r2 = st.columns(2)
with col_r1:
    numero_camion_saisi = st.text_input(
        "N° Camion / Tracteur / Remorque *",
        placeholder="Ex: CA7583RB ou CA7583RB/AP2290RB"
    )
with col_r2:
    masse_saisie = st.text_input(
        "Masse sur ticket (MT ou KG)",
        placeholder="Ex: 49.70 MT ou 49700 KG"
    )

resultats = []
message_alerte = ""

if numero_camion_saisi:
    resultats, message_alerte = rechercher_par_camion(
        numero_camion_saisi, masse_saisie, df_app, df_tracker, df_dech
    )
    if message_alerte:
        if "⚠️" in message_alerte:
            st.warning(message_alerte)
        else:
            st.info(message_alerte)

    if not resultats:
        st.error("Aucune correspondance. Ne pas enregistrer sans vérifier.")
    else:
        st.success(f"**{len(resultats)} résultat(s) trouvé(s)**")
        for i, r in enumerate(resultats):
            badge = ""
            if r.get("deja_decharge"):
                badge = " 🔴 DÉJÀ DÉCHARGÉ"
            elif r.get("priorite") == 1:
                badge = " 🟢 À DÉCHARGER"
            if r.get("score_masse") == "OK":
                badge += " ✅ Masse OK"
            elif r.get("score_masse") and "ÉCART" in str(r.get("score_masse", "")):
                badge += f" ⚠️ {r['score_masse']}"
            pu_aff = r.get("PU_FCFA_KG") or 0
            if pu_aff:
                badge += f" · PU {pu_aff:.0f} F/kg"

            with st.expander(
                f"📄 {r['Reference']} — {r['Fournisseur']} ({r['source']}){badge}",
                expanded=(i == 0 and r.get("priorite", 9) <= 1)
            ):
                st.write(f"**Référence BC :** `{r['Reference']}`")
                st.write(f"**Code FRS :** `{r.get('Code_Frs') or '—'}`")
                st.write(f"**Fournisseur :** {r['Fournisseur']}")
                st.write(f"**PU fournisseur :** {pu_aff:.0f} FCFA/kg" if pu_aff else "**PU :** non renseigné dans fournisseurs.xlsx")
                st.write(f"**Date :** {r['Date']}")
                st.write(f"**Quantité / Masse :** {r['Quantite']}")
                st.write(f"**Camion :** {r['Camion']}")
                if r.get("deja_decharge"):
                    st.error("Déjà déchargé — n'enregistrez que si nouveau passage.")

# ---------- Étape 2 ----------
st.markdown("---")
st.subheader("2️⃣ Enregistrer le déchargement (Ticket de Pont)")

bc_selectionne = ""
fournisseur_selectionne = ""
code_frs_sel = ""
pu_sel = 0.0
camion_selectionne = numero_camion_saisi if numero_camion_saisi else ""

if not resultats:
    st.info("Commencez par rechercher un N° Camion ci-dessus.")
else:
    options_bc = []
    for r in resultats:
        prefix = ""
        if r.get("priorite") == 1 and not r.get("deja_decharge"):
            prefix = "[À DÉCHARGER] "
        elif r.get("deja_decharge"):
            prefix = "[DÉJÀ FAIT] "
        options_bc.append(f"{prefix}{r['Reference']} | {r['Fournisseur']}")

    choix = st.selectbox("Sélectionnez le bon correspondant :", options_bc)
    idx = options_bc.index(choix)
    bc_selectionne = resultats[idx]["Reference"]
    fournisseur_selectionne = resultats[idx]["Fournisseur"]
    camion_selectionne = resultats[idx]["Camion"]
    code_frs_sel = resultats[idx].get("Code_Frs") or ""
    pu_sel = float(resultats[idx].get("PU_FCFA_KG") or 0)
    if pu_sel <= 0:
        pu_sel = charger_pu_fournisseur(code_frs_sel, fournisseur_selectionne)

    if resultats[idx].get("deja_decharge"):
        st.warning("Déchargement déjà enregistré — vérifiez avant de valider.")

with st.form("form_dechargement", clear_on_submit=True):
    col1, col2, col3 = st.columns(3)

    with col1:
        date_dechargement = st.date_input("Date de déchargement *", value=datetime.now())
        n_bc = st.text_input("N° BC *", value=bc_selectionne)
        code_frs = st.text_input("Code Frs", value=code_frs_sel)
        fournisseur = st.text_input("Fournisseur *", value=fournisseur_selectionne)

    with col2:
        n_camion = st.text_input("N° Camion *", value=camion_selectionne)
        weighment = st.text_input("Weighment Ticket (N° ticket pont) *")
        poids_net = st.number_input("Poids Net (KG) *", min_value=0.0, step=10.0)
        nb_sacs = st.number_input("Nb Sacs", min_value=0, step=1)

    with col3:
        pu = st.number_input(
            "PU (FCFA/KG) *",
            min_value=0.0,
            value=float(pu_sel) if pu_sel > 0 else 0.0,
            step=1.0,
            help="Prérempli depuis fournisseurs.xlsx — modifiable si besoin"
        )
        montant_paye = st.number_input("Montant Payé (FCFA)", min_value=0.0, step=1000.0)
        statut_paiement = st.selectbox("Statut Paiement", ["NON PAYE", "PARTIEL", "PAYE"])
        bon_paiement = st.text_input("Bon / N° Chèque de paiement")
        lieu_stockage = st.text_input("Lieu Stockage", value="GDIZ")

    montant_du = poids_net * pu if poids_net and pu else 0
    solde = montant_du - montant_paye
    st.markdown(
        f"**Calcul :** `{poids_net:,.0f} kg × {pu:,.0f} FCFA/kg = "
        f"**{montant_du:,.0f} FCFA**`  |  Solde : **{solde:,.0f} FCFA**"
    )

    submitted = st.form_submit_button("💾 Enregistrer le déchargement", type="primary", width="stretch")

    if submitted:
        if not n_bc.strip() or not n_camion.strip() or not weighment.strip() or poids_net <= 0:
            st.error("Champs obligatoires : N° BC, N° Camion, Ticket pont, Poids Net.")
        elif pu <= 0:
            st.error("PU obligatoire. Renseignez-le dans Fournisseurs ou saisissez-le ici.")
        else:
            df_exist = charger_dechargements()
            doublon = False
            if not df_exist.empty:
                mask = (
                    (df_exist["N_BC"].astype(str) == n_bc.strip()) &
                    (df_exist["N_Camion"].astype(str).str.upper().str.contains(
                        normaliser_camion(n_camion)[:8], na=False)) &
                    (df_exist["Weighment_Ticket"].astype(str) == weighment.strip())
                )
                if mask.any():
                    doublon = True

            if doublon:
                st.error("⚠️ Même BC + camion + ticket déjà enregistré.")
            else:
                ligne = {
                    "Date_Dechargement": date_dechargement.strftime("%d/%m/%Y"),
                    "N_BC": n_bc.strip(),
                    "Code_Frs": code_frs.strip(),
                    "Fournisseur": fournisseur.strip(),
                    "N_Camion": n_camion.strip(),
                    "Weighment_Ticket": weighment.strip(),
                    "Poids_Net_KG": poids_net,
                    "Nb_Sacs": nb_sacs,
                    "PU_FCFA_KG": pu,
                    "Montant_Du": montant_du,
                    "Montant_Paye": montant_paye,
                    "Solde_Restant": solde,
                    "Statut_Paiement": statut_paiement,
                    "Bon_Paiement": bon_paiement.strip(),
                    "Statut_Livraison": "DECHARGE",
                    "Lieu_Stockage": lieu_stockage.strip(),
                    "Utilisateur": name
                }
                ok, message = sauvegarder_dechargement(ligne)
                if ok:
                    st.success(f"✅ {message}")
                    st.success(
                        f"Bon **{n_bc}** — Camion **{n_camion}** — "
                        f"{poids_net:,.0f} kg × {pu:,.0f} = **{montant_du:,.0f} FCFA**"
                    )
                else:
                    st.error(message)
                st.rerun()

# ---------- Liste ----------
st.markdown("---")
st.subheader("📋 Déchargements déjà enregistrés")

df_dech = charger_dechargements()
if df_dech.empty:
    st.info("Aucun déchargement enregistré.")
else:
    col_f1, col_f2 = st.columns(2)
    with col_f1:
        filtre_statut = st.selectbox("Filtrer statut paiement", ["Tous", "NON PAYE", "PARTIEL", "PAYE"])
    with col_f2:
        filtre_cam = st.text_input("Filtrer N° camion", "")

    df_aff = df_dech.copy()
    if filtre_statut != "Tous" and "Statut_Paiement" in df_aff.columns:
        df_aff = df_aff[df_aff["Statut_Paiement"].astype(str).str.upper() == filtre_statut]
    if filtre_cam and "N_Camion" in df_aff.columns:
        df_aff = df_aff[df_aff["N_Camion"].astype(str).str.contains(filtre_cam, case=False, na=False)]

    st.dataframe(df_aff, width="stretch", height=300)
    with open(FICHIER_DECHARGEMENTS, "rb") as f:
        st.download_button(
            label="⬇️ Télécharger déchargements.xlsx",
            data=f,
            file_name="dechargements.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

st.markdown("---")
st.caption("NAP SARL – Montant = Poids × PU fournisseur · Anti-doublon · Priorité impayés")